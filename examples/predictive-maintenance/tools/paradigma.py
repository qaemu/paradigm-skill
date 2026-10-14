#!/usr/bin/env python3
"""PARADIGMA analysis engine.

Subcommands
  cpi     Download US CPI-U (FRED CPIAUCSL), save annual averages to cpi.csv
  curves  Robust cost/performance curves from points.csv (+ constraints.csv, cpi.csv)
  gap     Gap matrix from curves.csv + activity.csv (+ constraints.csv)
  papers  Yearly paper counts from OpenAlex for each query in queries.csv
  workbook Readable Excel workbook (Start here, Key findings, every table with sources)

Run `python3 paradigma.py <cmd> -h` for options. Only pandas, numpy, matplotlib
are required; scipy is used when present; openpyxl is needed for `workbook`.
"""
import argparse
import csv
import io
import json
import math
import os
import re
import sys
import urllib.parse
import urllib.request

import numpy as np
import pandas as pd

try:
    from scipy.stats import theilslopes as _theil
except Exception:  # scipy missing
    _theil = None


# ----------------------------------------------------------------- helpers
def theil_sen(x, y):
    """Return slope, intercept, slope_lo, slope_hi (95%)."""
    x = np.asarray(x, float)
    y = np.asarray(y, float)
    if _theil is not None:
        s, b, lo, hi = _theil(y, x, 0.95)
        return s, b, lo, hi
    slopes = [(y[j] - y[i]) / (x[j] - x[i])
              for i in range(len(x)) for j in range(i + 1, len(x)) if x[j] != x[i]]
    s = float(np.median(slopes))
    b = float(np.median(y - s * x))
    lo, hi = np.percentile(slopes, [2.5, 97.5])
    return s, b, float(lo), float(hi)


def parse_date(d):
    """'2014' | '2014-Q3' | '2014-07' | '2014-07-15' -> (decimal_year, year, quarter)."""
    d = str(d).strip()
    m = re.fullmatch(r"(\d{4})-?Q([1-4])", d, re.I)
    if m:
        y, q = int(m.group(1)), int(m.group(2))
        return y + (q - 0.5) / 4, y, q
    m = re.fullmatch(r"(\d{4})-(\d{1,2})(?:-(\d{1,2}))?", d)
    if m:
        y, mo = int(m.group(1)), int(m.group(2))
        return y + (mo - 0.5) / 12, y, (mo - 1) // 3 + 1
    m = re.fullmatch(r"(\d{4})(?:\.0)?", d)
    if m:
        y = int(m.group(1))
        return y + 0.5, y, None
    raise ValueError(f"bad date: {d!r}")


def pct_rank(s, ascending=True):
    s = pd.Series(s, dtype=float)
    r = s.rank(pct=True, ascending=ascending, method="average")
    return r


# --------------------------------------------------------------------- cpi
def cmd_cpi(a):
    url = "https://fred.stlouisfed.org/graph/fredgraph.csv?id=CPIAUCSL"
    raw = urllib.request.urlopen(url, timeout=30).read().decode()
    df = pd.read_csv(io.StringIO(raw))
    df.columns = ["date", "cpi"]
    df["cpi"] = pd.to_numeric(df["cpi"], errors="coerce")
    df["year"] = pd.to_datetime(df["date"]).dt.year
    g = df.dropna().groupby("year").agg(cpi=("cpi", "mean"), months=("cpi", "size")).reset_index()
    g["source"] = url
    g.to_csv(a.out, index=False)
    print(f"wrote {a.out}: {g.year.min()}-{g.year.max()} (last year months={int(g.months.iloc[-1])})")


# ------------------------------------------------------------------ curves
def cmd_curves(a):
    pts = pd.read_csv(a.points, dtype=str).fillna("")
    cons = pd.read_csv(a.constraints, dtype=str).fillna("") if os.path.exists(a.constraints) else pd.DataFrame()
    better = dict(zip(cons.get("constraint_id", []), cons.get("better", []))) if len(cons) else {}
    cpi = pd.read_csv(a.cpi) if a.cpi and os.path.exists(a.cpi) else None
    if cpi is not None:
        # a year counts as complete with >=10 monthly readings (e.g. Oct-2025 CPI was never published)
        full = cpi[cpi["months"] >= 10] if "months" in cpi else cpi
        base_year = int(a.base_year or full.year.max())
        cpi_map = dict(zip(cpi.year.astype(int), cpi.cpi.astype(float)))
        base_cpi = cpi_map[base_year]
    else:
        base_year, cpi_map, base_cpi = None, {}, None

    rows, rejects = [], []
    for i, r in pts.iterrows():
        why = None
        try:
            t, y, q = parse_date(r["date"])
            v = float(r["value"])
            if v <= 0:
                why = "value must be > 0 (log scale)"
        except Exception as e:
            why = str(e)
        if not why and r.get("source_url", "").strip() == "":
            why = "missing source_url"
        cur = (r.get("currency", "") or "").upper()
        basis = (r.get("price_basis", "") or "").lower()
        if not why and cur not in ("", "USD", "NA", "N/A"):
            why = f"currency {cur}: convert to USD before analysis"
        if not why and ((a.start and y < a.start) or (a.end and y > a.end)):
            why = f"outside window {a.start or ''}-{a.end or ''}"
        if why:
            rejects.append({**r.to_dict(), "reject_reason": why})
            continue
        real = v
        if cur == "USD" and basis.startswith("nominal"):
            if cpi_map.get(y) is None:
                rejects.append({**r.to_dict(), "reject_reason": f"no CPI for {y}"})
                continue
            real = v * base_cpi / cpi_map[y]
        elif cur == "USD" and basis.startswith("real-"):
            ry = int(basis.split("-")[1])
            if cpi_map.get(ry) is not None:
                real = v * base_cpi / cpi_map[ry]
        rows.append({"constraint_id": r["constraint_id"], "tier": r.get("tier", "") or "all",
                     "t": t, "year": y, "quarter": q, "value_real": real,
                     "source_url": r["source_url"].strip()})
    d = pd.DataFrame(rows)
    if rejects:
        pd.DataFrame(rejects).to_csv(os.path.join(a.outdir, "rejected_points.csv"), index=False)
    if d.empty:
        sys.exit("no usable points")

    per_rows, curve_rows = [], []
    for (cid, tier), g in d.groupby(["constraint_id", "tier"]):
        # choose period
        period = a.period
        if period == "auto":
            gq = g.dropna(subset=["quarter"])
            if len(gq) == len(g) and len(g) > 0:
                n_q = gq.groupby(["year", "quarter"]).size()
                period = "quarter" if (n_q >= a.min_n).mean() >= 0.6 and len(n_q) >= 8 else "year"
            else:
                period = "year"
        if period == "quarter":
            g = g.dropna(subset=["quarter"]).copy()
            g["period"] = g.year.astype(str) + "-Q" + g.quarter.astype(int).astype(str)
            g["pt"] = g.year + (g.quarter.astype(float) - 0.5) / 4
        else:
            g = g.copy()
            g["period"] = g.year.astype(str)
            g["pt"] = g.year + 0.5
        agg = g.groupby(["period", "pt"]).agg(
            n=("value_real", "size"), median=("value_real", "median"),
            q1=("value_real", lambda s: s.quantile(.25)), q3=("value_real", lambda s: s.quantile(.75)),
            sources=("source_url", "nunique")).reset_index().sort_values("pt")
        agg["low_confidence"] = agg.n < a.min_n
        # rolling median over 4 quarters to smooth sparse quarterly data
        if period == "quarter":
            agg["rolling4_median"] = agg["median"].rolling(4, min_periods=2).median()
        for _, p in agg.iterrows():
            per_rows.append({"constraint_id": cid, "tier": tier, "period_type": period, **p.to_dict()})

        # fit on period medians (one vote per period); fallback to points
        fit_on = agg if len(agg) >= 4 else g.drop(columns=["pt"]).rename(columns={"t": "pt", "value_real": "median"})
        x = fit_on["pt"].astype(float).values
        yv = np.log10(fit_on["median"].astype(float).values)
        res = {"constraint_id": cid, "tier": tier, "period_type": period,
               "n_points": len(g), "n_periods": len(agg), "n_sources": g.source_url.nunique(),
               "first_period": agg.period.iloc[0], "last_period": agg.period.iloc[-1],
               "low_conf_period_share": round(float(agg.low_confidence.mean()), 2),
               "better": better.get(cid, "lower") or "lower"}
        if len(np.unique(x)) < 3:
            res.update({"fit": "insufficient (<3 distinct periods)", "grade": "D"})
            curve_rows.append(res)
            continue
        s, b, lo, hi = theil_sen(x, yv)
        span = x.max() - x.min()
        fold = 10 ** (s * span)
        res.update({
            "fit": "theil-sen on log10(period medians)" if fit_on is agg else "theil-sen on log10(points)",
            "annual_change_pct": round((10 ** s - 1) * 100, 2),
            "annual_change_ci_low_pct": round((10 ** lo - 1) * 100, 2),
            "annual_change_ci_high_pct": round((10 ** hi - 1) * 100, 2),
            "fitted_start": round(10 ** (b + s * x.min()), 6),
            "fitted_end": round(10 ** (b + s * x.max()), 6),
            "fold_change_over_window": round(fold, 4),
            "span_years": round(span, 2),
        })
        # recent trend (last 5 years of data)
        recent = fit_on[fit_on.pt >= x.max() - 5]
        if recent.pt.nunique() >= 3:
            rs, *_ = theil_sen(recent.pt.astype(float).values, np.log10(recent["median"].astype(float).values))
            res["recent5y_annual_change_pct"] = round((10 ** rs - 1) * 100, 2)
        # improvement direction (positive = getting better)
        sign = -1 if res["better"] == "lower" else 1
        imp_lo, imp_hi = sorted([sign * lo, sign * hi])
        imp = sign * s
        improvement_fold = 10 ** (imp * span)
        res["improvement_fold"] = round(improvement_fold, 3)
        if improvement_fold >= 10:
            cls = "improved >=10x"
        elif imp_lo > 0 and (10 ** imp - 1) >= 0.05:
            cls = "improving"
        elif imp_hi < 0 and (10 ** (-imp) - 1) >= 0.05:
            cls = "worsening (constraint tightening)"
        else:
            cls = "flat / unclear"
        res["class"] = cls
        if "recent5y_annual_change_pct" in res:
            r5 = sign * math.log10(1 + res["recent5y_annual_change_pct"] / 100)
            res["trend_shape"] = ("accelerating" if r5 > imp * 1.25 else
                                  "decelerating" if r5 < imp * 0.75 else "steady")
        ci_w = abs(hi - lo) / max(abs(s), 1e-9)
        if len(agg) >= 8 and res["n_sources"] >= 3 and res["low_conf_period_share"] <= 0.2 and ci_w <= 1:
            grade = "A"
        elif len(agg) >= 5 and res["n_sources"] >= 2:
            grade = "B"
        else:
            grade = "C"
        res["grade"] = grade
        curve_rows.append(res)

    pd.DataFrame(per_rows).to_csv(os.path.join(a.outdir, "periods.csv"), index=False)
    cv = pd.DataFrame(curve_rows)
    cv.to_csv(os.path.join(a.outdir, "curves.csv"), index=False)
    meta = {"base_year": base_year, "min_n": a.min_n, "rejected": len(rejects),
            "window_start": a.start, "window_end": a.end}
    json.dump(meta, open(os.path.join(a.outdir, "curves_meta.json"), "w"), indent=2)
    print(cv.to_string(index=False))
    print(f"\nrejected points: {len(rejects)}  base year: {base_year}")
    if a.plots:
        plot_curves(d, pd.DataFrame(per_rows), cv, cons, a.outdir, base_year)


# ------------------------------------------------------------------- style
# One look for every chart: white page, blush plot area, near-black ink, one pink accent.
INK, MUTED, DEEP, PINK, ROSE, BLUSH, GRIDC = "#111111", "#55555E", "#C9406B", "#EE9AB3", "#F7CFDB", "#FFF6F8", "#E9C9D3"
WHYNOT_STYLE = {"BLOCKED": (DEEP, "Blocked: something still stands in the way"),
                "UNTRIED": (INK, "Untried: nobody has seriously built on it yet"),
                "TRIED TOO EARLY": (PINK, "Tried too early: failed before the constraint fell"),
                "N/A-CROWDED": (ROSE, "Crowded: many serious builders already"),
                "": ("#C4C4CC", "Not checked")}


def _mpl():
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    plt.rcParams.update({"font.size": 10, "axes.edgecolor": "#BBBBC2", "axes.labelcolor": INK,
                         "xtick.color": MUTED, "ytick.color": MUTED, "axes.titleweight": "bold"})
    return plt


def plot_curves(d, per, cv, cons, outdir, base_year):
    plt = _mpl()
    from matplotlib.ticker import MaxNLocator
    names = dict(zip(cons.get("constraint_id", []), cons.get("name", []))) if len(cons) else {}
    units = dict(zip(cons.get("constraint_id", []), cons.get("unit", []))) if len(cons) else {}
    os.makedirs(os.path.join(outdir, "plots"), exist_ok=True)
    for (cid, tier), g in d.groupby(["constraint_id", "tier"]):
        fig, ax = plt.subplots(figsize=(8, 4.6), facecolor="white")
        ax.set_facecolor(BLUSH)
        ax.grid(color=GRIDC, lw=0.8, which="major")
        ax.grid(color=GRIDC, lw=0.4, which="minor", alpha=.6)
        ax.scatter(g.t, g.value_real, s=16, color="#9A9AA3", alpha=.75, label="data points", zorder=2)
        p = per[(per.constraint_id == cid) & (per.tier == tier)]
        ax.plot(p.pt, p["median"], "o-", color=DEEP, lw=2, ms=5, mfc="white", mew=1.6, label="yearly median", zorder=3)
        c = cv[(cv.constraint_id == cid) & (cv.tier == tier)]
        sub = ""
        if len(c):
            r = c.iloc[0]
            if not pd.isna(r.get("fitted_start", np.nan)):
                x0, x1 = p.pt.min(), p.pt.max()
                ax.plot([x0, x1], [r.fitted_start, r.fitted_end], "--", color=INK, lw=1.4,
                        label=f"robust trend {r.annual_change_pct:+.1f}%/yr", zorder=4)
                sub = (f"{r.annual_change_pct:+.1f}% per year (95% range {r.annual_change_ci_low_pct:+.1f} to "
                       f"{r.annual_change_ci_high_pct:+.1f}) · grade {r.grade} · {int(r.n_points)} points, {int(r.n_sources)} sources")
            else:
                sub = f"too few years for a trend · grade {r.get('grade', 'D')} · {int(r.n_points)} points"
        ax.set_yscale("log")
        from matplotlib.ticker import FuncFormatter, LogLocator, NullFormatter
        ax.yaxis.set_major_locator(LogLocator(base=10, subs=(1.0, 2.0, 5.0)))
        ax.yaxis.set_major_formatter(FuncFormatter(lambda v, _: f"{v:,.6g}"))
        ax.yaxis.set_minor_formatter(NullFormatter())
        ax.xaxis.set_major_locator(MaxNLocator(integer=True))
        unit = units.get(cid, "")
        money = "USD" in str(unit).upper() or "$" in str(unit)
        ax.set_ylabel(f"{unit} ({base_year} dollars)" if (base_year and money) else unit)
        for sp in ("top", "right"):
            ax.spines[sp].set_visible(False)
        ax.legend(fontsize=8.5, frameon=False, loc="best")
        fig.text(0.01, 0.97, f"{cid} {names.get(cid, cid)}", fontsize=12.5, fontweight="bold", color=INK, va="top")
        fig.text(0.01, 0.905, f"series: {tier}" + (f" · {sub}" if sub else ""), fontsize=8.8, color=MUTED, va="top")
        fig.subplots_adjust(left=0.1, right=0.98, top=0.84, bottom=0.1)
        fig.savefig(os.path.join(outdir, "plots", f"curve_{cid}_{tier}.png"), dpi=150,
                    bbox_inches="tight", pad_inches=0.25, facecolor="white")
        plt.close(fig)


def plot_gap(df, cut, outdir, title, whynot):
    """Gap matrix with ID labels on the dots and a key beside the chart. Nothing is clipped."""
    import textwrap
    plt = _mpl()
    s = df.dropna(subset=["drop_score", "activity_score"]).copy()
    lab = {cid: str(whynot.get(cid, "")).strip().upper() for cid in s.constraint_id}
    n = len(s)
    rows_per_col = 16
    ncols = 1 if n <= rows_per_col else 2
    key_w = 0.37 if ncols == 1 else 0.52
    fig_w = 13 if ncols == 1 else 16
    fig_h = max(7.6, 0.5 + 0.42 * min(n, rows_per_col) + 3.2)
    fig = plt.figure(figsize=(fig_w, fig_h), facecolor="white")
    ax_w = 0.95 - key_w - 0.1
    ax = fig.add_axes([0.07, 0.13, ax_w, 0.74])
    ax.set_facecolor(BLUSH)
    ax.axvline(cut, color=GRIDC, lw=1.3)
    ax.axhline(cut, color=GRIDC, lw=1.3)
    for t, (x, yt, ys, ha, col, sub) in {
            "OPEN GAP": (.02, .97, .925, "left", DEEP, "dropped a lot, few use it yet"),
            "CROWDED": (.98, .97, .925, "right", MUTED, "dropped a lot, many already use it"),
            "DORMANT": (.02, .075, .03, "left", MUTED, "barely moved, little activity"),
            "HYPE RISK": (.98, .075, .03, "right", MUTED, "barely moved, many building anyway")}.items():
        va = "top" if yt > .5 else "bottom"
        ax.text(x, yt, t, transform=ax.transAxes, ha=ha, va=va, fontsize=11, fontweight="bold", color=col)
        ax.text(x, ys, sub, transform=ax.transAxes, ha=ha, va=va, fontsize=8.5, color=MUTED)
    # spread dots that sit on the same spot so every label stays readable
    s["px"], s["py"] = s.activity_score.astype(float), s.drop_score.astype(float)
    for _, grp in s.groupby([s.px.round(2), s.py.round(2)]):
        if len(grp) > 1:
            offs = np.linspace(-0.025 * (len(grp) - 1), 0.025 * (len(grp) - 1), len(grp))
            s.loc[grp.index, "px"] = grp.px + offs
    for _, r in s.iterrows():
        col = WHYNOT_STYLE.get(lab[r.constraint_id], WHYNOT_STYLE[""])[0]
        ax.scatter(r.px, r.py, s=190, color=col, edgecolor=INK, linewidth=1.2, zorder=3)
        ax.annotate(r.constraint_id, (r.px, r.py), xytext=(0, 12), textcoords="offset points",
                    ha="center", fontsize=10, fontweight="bold", color=INK)
    ax.set_xlim(-.06, 1.06)
    ax.set_ylim(-.06, 1.24)
    ax.set_yticks([0, .2, .4, .6, .8, 1.0])
    ax.set_xlabel("How many are building on it  →\n(percentile of companies, funding and launches)", fontsize=10)
    ax.set_ylabel("How far the constraint fell  →\n(percentile)", fontsize=10)
    fig.text(0.07, 0.95, title, fontsize=14, fontweight="bold", color=INK)
    # key
    k = fig.add_axes([0.07 + ax_w + 0.04, 0.05, key_w, 0.86])
    k.axis("off"); k.set_xlim(0, 1); k.set_ylim(0, 1)
    k.text(0, 0.985, "Constraints on the chart", fontsize=11.5, fontweight="bold", color=INK, va="top")
    step = min(0.052, 0.58 / max(1, min(n, rows_per_col)))
    colw = 1.0 / ncols
    namew = 34 if ncols == 1 else 30
    for i, (_, r) in enumerate(s.sort_values("constraint_id").iterrows()):
        cx = (i // rows_per_col) * colw
        y = 0.92 - (i % rows_per_col) * step
        col = WHYNOT_STYLE.get(lab[r.constraint_id], WHYNOT_STYLE[""])[0]
        k.scatter([cx + 0.02], [y], s=100, color=col, edgecolor=INK, linewidth=1, clip_on=False)
        k.text(cx + 0.07 * (1 if ncols == 1 else 0.75), y, r.constraint_id, fontsize=9.5, fontweight="bold", color=INK, va="center")
        nm = str(r.get("short_name", "") or "").strip()
        nm = nm if nm and nm != "nan" else str(r["name"])
        nm = nm if len(nm) <= namew else nm[:namew - 1].rstrip() + "…"
        k.text(cx + 0.18 * (1 if ncols == 1 else 0.6), y, nm, fontsize=9.5, color=INK, va="center")
    y = 0.92 - min(n, rows_per_col) * step - 0.03
    k.text(0, y, "Dot colour = result of the why-not check", fontsize=9.5, fontweight="bold", color=INK, va="center")
    y -= 0.045
    for key in ("BLOCKED", "UNTRIED", "TRIED TOO EARLY", "N/A-CROWDED", ""):
        col, txt = WHYNOT_STYLE[key]
        k.scatter([0.02], [y], s=80, color=col, edgecolor=INK, linewidth=1, clip_on=False)
        k.text(0.07 if ncols == 1 else 0.05, y, txt, fontsize=9, color=INK, va="center")
        y -= 0.038
    off = df[df.drop_score.isna() | df.activity_score.isna()]
    if len(off):
        parts = []
        for _, r in off.iterrows():
            why = str(r.get("not_placed_reason", "") or "").strip()
            if not why or why == "nan":
                why = "no measured change" if pd.isna(r.drop_score) else "no activity data"
            parts.append(f"{r.constraint_id} ({why})")
        note = "Not on the chart: " + "; ".join(parts) + "."
        wrap = 66 if ncols == 1 else 95
        k.text(0, y - 0.02, "\n".join(textwrap.wrap(note, wrap)), fontsize=8.6, color=MUTED, va="top", linespacing=1.35)
    os.makedirs(os.path.join(outdir, "plots"), exist_ok=True)
    fig.savefig(os.path.join(outdir, "plots", "gap_matrix.png"), dpi=150, bbox_inches="tight", pad_inches=0.35, facecolor="white")
    plt.close(fig)


# --------------------------------------------------------------------- gap
def cmd_gap(a):
    cons_all = pd.read_csv(a.constraints, dtype=str).fillna("")
    act = pd.read_csv(a.activity)
    cv = pd.read_csv(a.curves) if os.path.exists(a.curves) else pd.DataFrame()
    # Only loosening constraints compete for open gap. Tightening, mixed or pressure constraints, and any row
    # with an exclude_reason, are kept out of the percentile ranking and listed under the chart instead.
    direction = cons_all.get("direction", pd.Series([""] * len(cons_all))).str.lower().str.strip()
    excl = cons_all.get("exclude_reason", pd.Series([""] * len(cons_all))).str.strip()
    out = direction.isin(["tightening", "mixed", "pressure"]) | (excl != "")
    reason = np.where(excl != "", excl, direction + ", not an enabler")
    cons = cons_all[~out]
    left_out = cons_all[out].assign(not_placed_reason=reason[out.values])
    df = cons[["constraint_id", "name", "family"] + (["short_name"] if "short_name" in cons else [])].copy()

    # improvement score: measured curves -> log10(improvement_fold) per decade; qualitative 0-3 otherwise
    imp = {}
    if len(cv) and "improvement_fold" in cv:
        best = cv.dropna(subset=["improvement_fold"]).sort_values("n_points", ascending=False)
        best = best.drop_duplicates("constraint_id")
        for _, r in best.iterrows():
            per_decade = math.log10(max(r.improvement_fold, 1e-9)) / max(r.span_years, 1) * 10
            imp[r.constraint_id] = (per_decade, r.get("grade", ""), "measured")
    qual = dict(zip(cons.constraint_id, cons.get("qualitative_change_0to3", [""] * len(cons))))
    vals, grades, kinds = [], [], []
    for cid in df.constraint_id:
        if cid in imp:
            v, gr, k = imp[cid]
        elif str(qual.get(cid, "")).strip() not in ("", "nan"):
            # 0..3 qualitative mapped onto log10-per-decade scale (3 ~ 10x/decade)
            v, gr, k = float(qual[cid]) / 3, "Q", "qualitative"
        else:
            v, gr, k = np.nan, "", "none"
        vals.append(v); grades.append(gr); kinds.append(k)
    df["improvement_log10_per_decade"] = vals
    df["improvement_basis"] = kinds
    df["curve_grade"] = grades
    df = df.merge(act, on="constraint_id", how="left")
    if "threshold_crossed" in df:
        df["threshold_crossed"] = df.threshold_crossed.astype(str).str.lower().isin(["1", "true", "yes", "y"])
    else:
        df["threshold_crossed"] = False

    df["drop_score"] = pct_rank(df.improvement_log10_per_decade)
    df.loc[df.threshold_crossed, "drop_score"] = np.minimum(1, df.loc[df.threshold_crossed, "drop_score"] + 0.15)

    signals = [c for c in ["companies_count", "funding_usd_5y", "patents_last3y", "launches_count", "job_postings"]
               if c in df]
    ranks = [pct_rank(np.log1p(pd.to_numeric(df[c], errors="coerce"))) for c in signals]
    df["activity_score"] = pd.concat(ranks, axis=1).mean(axis=1, skipna=True) if ranks else np.nan
    df["activity_signals_used"] = df[signals].notna().sum(axis=1) if signals else 0

    if "papers_last3y" in df:
        papers = pd.to_numeric(df.papers_last3y, errors="coerce")
        comp = pd.to_numeric(df.get("companies_count", 0), errors="coerce").fillna(0)
        df["research_to_commerce"] = (papers / (comp + 1)).round(1)
        df["research_ahead_score"] = pct_rank(np.log1p(df.research_to_commerce))
    if "papers_prev3y" in df and "papers_last3y" in df:
        df["papers_growth"] = (pd.to_numeric(df.papers_last3y, errors="coerce") /
                               pd.to_numeric(df.papers_prev3y, errors="coerce").replace(0, np.nan)).round(2)

    hi_d = df.drop_score >= a.cut
    hi_a = df.activity_score >= a.cut
    df["quadrant"] = np.select(
        [hi_d & ~hi_a, hi_d & hi_a, ~hi_d & hi_a, ~hi_d & ~hi_a],
        ["OPEN GAP (fell, few builders)", "CROWDED (fell, many builders)",
         "HYPE RISK (little change, many builders)", "DORMANT"], default="UNSCORED")
    df.loc[df.drop_score.isna() | df.activity_score.isna(), "quadrant"] = "UNSCORED"
    df["gap_score"] = (df.drop_score * (1 - df.activity_score)).round(3)
    if "research_ahead_score" in df:
        df["gap_score"] = (df.gap_score * (0.8 + 0.4 * df.research_ahead_score.fillna(0.5))).round(3)
    df = df.sort_values("gap_score", ascending=False)
    if len(left_out):
        lo = left_out[["constraint_id", "name", "family", "not_placed_reason"]].merge(act, on="constraint_id", how="left")
        lo["quadrant"] = "NOT PLACED"
        df = pd.concat([df, lo], ignore_index=True)
    df.to_csv(os.path.join(a.outdir, "gap_matrix.csv"), index=False)
    cols = ["constraint_id", "name", "drop_score", "activity_score", "gap_score", "quadrant"]
    print(df[[c for c in cols if c in df]].to_string(index=False))

    if a.plots:
        wn = {}
        if a.whynot and os.path.exists(a.whynot):
            w = pd.read_csv(a.whynot, dtype=str).fillna("")
            wn = dict(zip(w.constraint_id, w.label))
        title = a.title or "Gap matrix"
        plot_gap(df, a.cut, a.outdir, title, wn)


# ------------------------------------------------------------------ papers
def cmd_papers(a):
    q = pd.read_csv(a.queries, dtype=str)
    key = os.environ.get("OPENALEX_API_KEY", "")
    out = []
    for _, r in q.iterrows():
        params = {"filter": f"title_and_abstract.search:{r['query']},publication_year:{a.start}-{a.end}",
                  "group_by": "publication_year"}
        if key:
            params["api_key"] = key
        url = "https://api.openalex.org/works?" + urllib.parse.urlencode(params)
        try:
            js = json.loads(urllib.request.urlopen(url, timeout=30).read().decode())
            for gr in js.get("group_by", []):
                out.append({"constraint_id": r["constraint_id"], "query": r["query"],
                            "year": int(gr["key"]), "papers": int(gr["count"]),
                            "source_url": url.split("&api_key")[0]})
            print(f"ok   {r['constraint_id']}: {r['query']}")
        except Exception as e:
            body = getattr(e, "read", lambda: b"")().decode(errors="ignore")[:160]
            print(f"FAIL {r['constraint_id']}: {e} {body}")
    if not out:
        sys.exit("no OpenAlex data (set OPENALEX_API_KEY or use the web fallback)")
    df = pd.DataFrame(out)
    df.to_csv(a.out, index=False)
    last = a.end
    l3 = df[df.year.between(last - 2, last)].groupby("constraint_id").papers.sum().rename("papers_last3y")
    p3 = df[df.year.between(last - 5, last - 3)].groupby("constraint_id").papers.sum().rename("papers_prev3y")
    summ = pd.concat([l3, p3], axis=1).fillna(0).astype(int)
    summ.to_csv(os.path.splitext(a.out)[0] + "_summary.csv")
    print(summ.to_string())


# ---------------------------------------------------------------- workbook
GLOSSARY = [
    ("Constraint", "Anything that limits what the field can do: a price, a performance level, a law, a patent, a skill shortage, a norm, a supply problem."),
    ("Fell / loosening", "The constraint got weaker (cheaper, faster, legal, allowed). Something may have just become possible."),
    ("Tightening", "The constraint got stronger (more expensive, more regulated, scarcer)."),
    ("Fold change", "How many times cheaper or better something got. 10× means it now costs a tenth."),
    ("Change per year (%)", "Typical yearly change, computed with a method that ignores outliers (Theil–Sen). The low and high values are the 95% range."),
    ("Real USD", "Prices adjusted for inflation to the latest full year's dollars, so years can be compared."),
    ("Grade A–D", "How much to trust a curve. A and B: enough points and sources. C: usable but weak. D: too little data, no trend computed."),
    ("Drop score", "0 to 1. How far a constraint fell compared with the others in this run."),
    ("Activity score", "0 to 1. How many companies, how much funding and how many launches build on it, compared with the others."),
    ("Open gap", "The constraint dropped a lot (cheaper, allowed, easier), but few companies use the change yet. A possible opening, but only after the why-not check."),
    ("Crowded", "The constraint dropped a lot and many companies already use it. This is what most people in the field already know."),
    ("Hype risk", "The constraint barely moved, yet many companies are building on it anyway. A warning sign: lots of activity with no real change underneath."),
    ("Dormant", "The constraint barely moved and few are building on it."),
    ("Not placed", "Left off the gap matrix on purpose: it tightened, it is a pressure rather than an enabler, or it could not be measured."),
    ("Untried", "Nobody has seriously built on it yet."),
    ("Tried too early", "Someone failed before the constraint fell. Could work now."),
    ("Blocked", "Something else still stands in the way. Goes to a watch list, not to open gap."),
    ("Company status", "announced = press release, funding round or pilot only. deployed = named customers or installations. paying = evidence of revenue."),
    ("Consensus", "yes = widely covered, well funded or a big incumbent. no = less known."),
    ("Verification", "Separate agents that did not collect the data re-opened each source. MATCH = the page shows it. MISMATCH = it does not (then corrected). UNVERIFIABLE = the page could not be opened."),
]

TAGS = {  # value -> (fill, text colour); dark text on light fills keeps contrast high
    "MATCH": ("DCF2E1", "14532D"), "MISMATCH": ("FAD7DC", "7F1D1D"), "UNVERIFIABLE": ("FFF0BF", "6B3A0A"),
    "loosening": ("DCF2E1", "14532D"), "tightening": ("FAD7DC", "7F1D1D"), "mixed": ("FFF0BF", "6B3A0A"), "pressure": ("FFF0BF", "6B3A0A"),
    "OPEN GAP": ("FBE0E8", "8A1C3F"), "CROWDED": ("E3E8F0", "1F2937"), "HYPE RISK": ("FFF0BF", "6B3A0A"),
    "DORMANT": ("ECECEF", "374151"), "UNSCORED": ("FFFFFF", "4B4B55"), "NOT PLACED": ("FFFFFF", "4B4B55"),
    "UNTRIED": ("DBE7FE", "1E3A8A"), "BLOCKED": ("FAD7DC", "7F1D1D"), "N/A-CROWDED": ("E3E8F0", "1F2937"),
    "TRIED TOO EARLY": ("FFF0BF", "6B3A0A"),
    "paying": ("DCF2E1", "14532D"), "deployed": ("DBE7FE", "1E3A8A"), "announced": ("ECECEF", "374151"),
    "A": ("DCF2E1", "14532D"), "B": ("DCF2E1", "14532D"), "C": ("FFF0BF", "6B3A0A"), "D": ("FAD7DC", "7F1D1D"),
    "High": ("DCF2E1", "14532D"), "Medium": ("DBE7FE", "1E3A8A"), "Low": ("FFF0BF", "6B3A0A"),
    "improved >=10x": ("DCF2E1", "14532D"), "improving": ("DCF2E1", "14532D"),
    "flat / unclear": ("ECECEF", "374151"), "worsening (constraint tightening)": ("FAD7DC", "7F1D1D"),
    "enacted": ("ECECEF", "374151"), "proposed": ("FFF0BF", "6B3A0A"), "expiring": ("FFF0BF", "6B3A0A"), "expired": ("ECECEF", "374151"),
}


def _tag(v):
    v = str(v).strip()
    if v in TAGS:
        return TAGS[v]
    for k in ("OPEN GAP", "CROWDED", "HYPE RISK", "DORMANT"):
        if v.startswith(k):
            return TAGS[k]
    if v[:1] in "ABCD" and v[1:2] in (" ", "("):
        return TAGS[v[0]]
    return None


def _lines(text, width):
    if text is None or str(text) == "":
        return 1
    per_line = max(4, width * 1.15)
    return sum(max(1, math.ceil(len(part) / per_line)) for part in str(text).split("\n"))


def cmd_workbook(a):
    """Readable Excel workbook for a reader with no background in market research."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        sys.exit("workbook needs openpyxl: pip install openpyxl")

    def R(f):
        return pd.read_csv(f, dtype=str).fillna("") if os.path.exists(f) else pd.DataFrame()

    F = "Arial"
    head_fill = PatternFill("solid", fgColor="1F2430")
    zebra = PatternFill("solid", fgColor="F3F4F7")
    side = Side(style="thin", color="9CA0AA")
    border = Border(left=side, right=side, top=side, bottom=side)
    f_body, f_head = Font(name=F, size=10, color="111111"), Font(name=F, size=10, bold=True, color="FFFFFF")
    f_link, f_bold = Font(name=F, size=10, color="1D4ED8", underline="single"), Font(name=F, size=10, bold=True, color="111111")
    f_muted = Font(name=F, size=10, color="4B4B55")

    def fit(ws, landscape=True):
        ws.page_setup.orientation = "landscape" if landscape else "portrait"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0

    def table(wb, name, df, cols, intro=None, tab=None):
        """cols: (source column, header, width, kind); kind in text|num|score|int|usd|pct|url|tag."""
        ws = wb.create_sheet(name[:31])
        if tab:
            ws.sheet_properties.tabColor = tab
        r0 = 1
        if intro:
            c = ws.cell(1, 1, intro)
            c.font = Font(name=F, size=10, italic=True, color="4B4B55")
            c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(cols))
            ws.row_dimensions[1].height = 15 * _lines(intro, sum(c[2] for c in cols)) + 6
            r0 = 3
        for j, (_, head, width, _) in enumerate(cols, 1):
            c = ws.cell(r0, j, head)
            c.font, c.fill, c.border = f_head, head_fill, border
            c.alignment = Alignment(wrap_text=True, vertical="center")
            ws.column_dimensions[get_column_letter(j)].width = width
        ws.row_dimensions[r0].height = 15 * max(_lines(h, w) for _, h, w, _ in cols) + 8
        if df is None or df.empty:
            ws.cell(r0 + 1, 1, "No data for this sheet in this run.").font = f_muted
            fit(ws)
            return ws
        for i, (_, row) in enumerate(df.iterrows(), 1):
            r, lines = r0 + i, 1
            for j, (src, _, width, kind) in enumerate(cols, 1):
                v = row[src] if src in row.index else ""
                c = ws.cell(r, j)
                c.border, c.font = border, f_body
                c.alignment = Alignment(wrap_text=True, vertical="top")
                if i % 2 == 0:
                    c.fill = zebra
                if v is None or str(v) in ("", "nan"):
                    continue
                if kind in ("num", "score", "int", "usd", "pct"):
                    try:
                        c.value = float(str(v).replace(",", ""))
                        c.number_format = {"num": "#,##0.00##", "score": "0.00", "int": "#,##0",
                                           "usd": "$#,##0", "pct": '0.0"%"'}[kind]
                        c.alignment = Alignment(vertical="top", horizontal="right")
                    except ValueError:
                        c.value = v
                elif kind == "url":
                    urls = [u for u in str(v).split() if u.startswith("http")]
                    if urls:
                        c.value = "Open source" + (f" (+{len(urls) - 1} more)" if len(urls) > 1 else "")
                        c.hyperlink, c.font = urls[0], f_link
                    else:
                        c.value = v
                else:
                    c.value = v
                    lines = max(lines, _lines(v, width))
                    if kind == "tag" and _tag(v):
                        fill, ink = _tag(v)
                        c.fill = PatternFill("solid", fgColor=fill)
                        c.font = Font(name=F, size=10, bold=True, color=ink)
            ws.row_dimensions[r].height = min(15 * lines + 4, 409)
        ws.auto_filter.ref = f"A{r0}:{get_column_letter(len(cols))}{r0 + len(df)}"
        ws.freeze_panes = ws.cell(r0 + 1, 2)
        fit(ws)
        return ws

    cons = R(a.constraints)
    names = dict(zip(cons.get("constraint_id", []), cons.get("name", [])))
    lab = lambda cid: f"{cid} {names.get(cid, '')}".strip()
    findings, gap, wn, cv, per = R("findings.csv"), R("gap_matrix.csv"), R("whynot.csv"), R("curves.csv"), R("periods.csv")
    pts, nc, co, fe, act = R(a.points), R("noncost.csv"), R("companies.csv"), R("funding_events.csv"), R("activity.csv")
    fail, absn, ver, review = R("failures.csv"), R("absence.csv"), R("verification.csv"), R("curve_review.csv")
    window = f"{a.start}–{a.end}" if a.start and a.end else "see report"

    wb = Workbook()
    ws = wb.active
    ws.title = "Start here"
    ws.sheet_properties.tabColor = "C9406B"
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 26, 100
    state = {"row": 1}

    def put(x, y="", style="pair"):
        r = state["row"]
        ca, cb = ws.cell(r, 1, x), ws.cell(r, 2, y)
        for c in (ca, cb):
            c.alignment, c.font = Alignment(wrap_text=True, vertical="top"), f_body
        if style == "title":
            ca.font = Font(name=F, size=16, bold=True, color="111111")
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
            ws.row_dimensions[r].height = 28
        elif style == "section":
            ca.font = Font(name=F, size=12, bold=True, color="FFFFFF")
            ca.fill = cb.fill = head_fill
            ca.alignment = Alignment(vertical="center")
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
            ws.row_dimensions[r].height = 22
        elif style == "note":
            ca.font = f_muted
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
            ws.row_dimensions[r].height = 15 * _lines(x, 126) + 6
        else:
            ca.font = f_bold
            ca.border = cb.border = border
            ws.row_dimensions[r].height = 15 * max(_lines(y, 100), _lines(x, 26)) + 6
        state["row"] += 1

    put(f"PARADIGMA: {a.field or 'field scan'}", style="title")
    put(f"Run date {a.date or ''} · region {a.region or 'global'} · window {window} · {len(cons)} constraints. "
        "This workbook holds every number behind the report, with a link to where it came from.", style="note")
    state["row"] += 1
    put("What this is", style="section")
    put("PARADIGMA asks one question about a field: which of the things that limit it (prices, performance, laws, patents, skills, "
        "norms, supply) changed over the chosen years, who is already building on each change, and where something changed but "
        "almost nobody is building yet. A limit that fell but nobody uses yet is a possible opening, called open gap.", style="note")
    state["row"] += 1
    put("Read it in 5 minutes", style="section")
    for x, y in [("1. Key findings", "The plain-language results. Start here."),
                 ("2. Gap matrix", "Each constraint scored on how far it fell and how many are building on it. The quadrant says what that combination means."),
                 ("3. Why-not check", "For each possible opening: did someone already try and fail, or is something still blocking it?"),
                 ("4. Constraints", "The limits that were studied and why each matters to the field."),
                 ("5. Everything else", "The evidence: price curves, data points, laws, companies, funding, and the independent verification of each claim.")]:
        put(x, y)
    state["row"] += 1
    put("Words used in this workbook", style="section")
    for x, y in GLOSSARY:
        put(x, y)
    state["row"] += 1
    put("Sheets", style="section")
    for x, y in [("Key findings", "Plain-language results with confidence."), ("Gap matrix", "Scores and quadrant per constraint."),
                 ("Why-not check", "Failed attempts and remaining blockers for each candidate."),
                 ("Constraints", "What each constraint measures and the change scores."), ("Price curves", "Trend per series: change per year, fold change, grade."),
                 ("Yearly medians", "The median value per year behind each curve."), ("Data points", "Every price or performance number, with its source."),
                 ("Laws & timeline", "Laws, standards, patents and supply events, sorted by date."), ("Companies", "Who builds on each constraint, with status and evidence."),
                 ("Funding rounds", "Disclosed rounds in the activity window."), ("Activity", "Counts per constraint: companies, funding, papers."),
                 ("Failures & exits", "Companies that shut down, retreated or sold."), ("Searches, no result", "Searches that found nothing relevant (evidence of absence)."),
                 ("Verification", "The independent re-check of each claim.")]:
        put(x, y)
    ws.freeze_panes = "A2"
    fit(ws, landscape=False)

    table(wb, "Key findings", findings, [
        ("finding", "Finding", 34, "text"), ("constraint_ids", "Constraint", 11, "text"), ("number", "The number", 46, "text"),
        ("confidence", "Confidence (grade)", 12, "tag"), ("meaning", "What it means", 62, "text")],
        intro="The results in plain words. Each row points to a constraint ID you can look up in the other sheets.", tab="C9406B")

    if not gap.empty:
        gap = gap.copy()
        gap["constraint"] = gap.constraint_id.map(lab)
        gap["quad"] = gap.quadrant.str.split(" \\(").str[0]
        gap["why"] = gap.constraint_id.map(dict(zip(wn.get("constraint_id", []), wn.get("label", []))))
    table(wb, "Gap matrix", gap, [
        ("constraint", "Constraint", 44, "text"), ("drop_score", "How far it fell (0–1)", 13, "score"),
        ("activity_score", "How many build on it (0–1)", 14, "score"), ("gap_score", "Opening score (higher = more open)", 15, "score"),
        ("quad", "Quadrant", 15, "tag"), ("why", "Why-not result", 15, "tag"), ("not_placed_reason", "Why not placed", 24, "text"),
        ("companies_count", "Companies found", 11, "int"), ("funding_usd_5y", "Funding to those companies", 16, "usd"),
        ("papers_last3y", "Research papers, last 3 years", 12, "int"), ("papers_prev3y", "Research papers, 3 years before", 12, "int")],
        intro="Quadrants come from the two 0–1 scores. UNSCORED means one of the two signals is missing. NOT PLACED means the constraint tightened, is a pressure, or could not be measured.", tab="1F2430")

    if not wn.empty:
        wn = wn.copy()
        wn["constraint"] = wn.constraint_id.map(lab)
    table(wb, "Why-not check", wn, [
        ("constraint", "Constraint", 30, "text"), ("label", "Result", 14, "tag"), ("failed_attempts", "Who tried and what happened", 60, "text"),
        ("blockers", "What still blocks it", 60, "text"), ("blocker_trend", "Are the blockers falling?", 36, "text"),
        ("source_urls", "Source", 16, "url"), ("note", "Notes", 50, "text")],
        intro="Only Untried and Tried too early count as open gap. Blocked goes to a watch list.")

    c2 = cons.copy()
    if not c2.empty:
        q = c2.get("qualitative_change_0to3", pd.Series([""] * len(c2)))
        c2["score"] = q
    table(wb, "Constraints", c2, [
        ("constraint_id", "ID", 7, "text"), ("name", "Constraint", 36, "text"), ("family", "Type", 11, "text"),
        ("unit", "Measured in", 18, "text"), ("better", "Better when", 9, "text"), ("why_it_matters", "Why it matters", 44, "text"),
        ("direction", "Direction", 13, "tag"), ("exclude_reason", "Left off the matrix because", 24, "text"),
        ("score", "Change score 0–3 (if not measured)", 13, "int"), ("qualitative_justification", "Reason for the score", 70, "text")],
        intro="0 = no change, 1 = minor, 2 = significant, 3 = a hard barrier was removed (or added, if tightening).")

    if not cv.empty:
        cv = cv.copy()
        cv["constraint"] = cv.constraint_id.map(lab)
        cv["years"] = cv.first_period.astype(str) + "–" + cv.last_period.astype(str)
        cv["review"] = ""
        if not review.empty:
            for _, rv in review.iterrows():
                m = (cv.constraint_id == rv.get("constraint_id", "")) & ((cv.tier == rv.get("tier", "")) | (rv.get("tier", "") == ""))
                if rv.get("grade_override", ""):
                    cv.loc[m, "grade"] = f"{rv.grade_override} (engine said " + cv.loc[m, "grade"].astype(str) + ")"
                cv.loc[m, "review"] = rv.get("note", "")
    table(wb, "Price curves", cv, [
        ("constraint", "Constraint", 32, "text"), ("tier", "Series", 20, "text"), ("years", "Years", 11, "text"),
        ("n_points", "Data points", 8, "int"), ("n_sources", "Sources", 8, "int"), ("annual_change_pct", "Change per year %", 11, "pct"),
        ("annual_change_ci_low_pct", "Range low %", 10, "pct"), ("annual_change_ci_high_pct", "Range high %", 10, "pct"),
        ("improvement_fold", "Times better over the period (1 = no change)", 14, "num"), ("class", "Verdict", 18, "tag"),
        ("trend_shape", "Speeding up or slowing?", 13, "text"), ("grade", "Grade", 12, "tag"), ("review", "Reviewer note", 46, "text")],
        intro="Negative change per year means the price fell. Times better below 1 means the constraint tightened. Grade D rows had too little data for a trend.")

    if not per.empty:
        per = per.copy()
        per["constraint"] = per.constraint_id.map(lab)
    table(wb, "Yearly medians", per, [
        ("constraint", "Constraint", 32, "text"), ("tier", "Series", 20, "text"), ("period", "Year", 8, "text"),
        ("n", "Points that year", 9, "int"), ("median", "Median (real USD or unit)", 14, "num"), ("q1", "Lower quartile", 12, "num"),
        ("q3", "Upper quartile", 12, "num"), ("sources", "Sources", 8, "int"), ("low_confidence", "Fewer than 3 points?", 11, "text")],
        intro="One row per series and year. Money values are inflation-adjusted.")

    if not pts.empty:
        pts = pts.copy()
        pts["constraint"] = pts.constraint_id.map(lab)
        if "row_id" not in pts:
            pts["row_id"] = [f"P{i:03d}" for i in range(len(pts))]
    table(wb, "Data points", pts, [
        ("row_id", "Row", 7, "text"), ("constraint", "Constraint", 28, "text"), ("tier", "Series", 18, "text"), ("date", "Date", 11, "text"),
        ("value", "Value as published", 12, "num"), ("unit", "Unit", 16, "text"), ("currency", "Currency", 8, "text"),
        ("price_basis", "Nominal or real", 9, "text"), ("source_title", "Source", 40, "text"), ("source_url", "Link", 14, "url"),
        ("source_type", "Source type", 10, "text"), ("note", "How the value was taken", 60, "text")],
        intro="Values exactly as published (not inflation-adjusted).")

    if not nc.empty:
        nc = nc.copy().sort_values("date_effective")
        nc["constraint"] = nc.constraint_id.map(lab)
    table(wb, "Laws & timeline", nc, [
        ("date_effective", "Date", 11, "text"), ("constraint", "Constraint", 28, "text"), ("change", "What changed", 62, "text"),
        ("direction", "Direction", 11, "tag"), ("status", "Status", 10, "tag"), ("jurisdiction", "Where", 12, "text"),
        ("source_url", "Link", 14, "url"), ("note", "Notes", 60, "text"), ("item_id", "ID", 7, "text")],
        intro=f"Sorted by date. Dates after {a.date or 'today'} are upcoming. Patent expiry dates are approximate.")

    if not co.empty:
        co = co.copy()
        co["constraint"] = co.constraint_id.map(lab)
        co = co.sort_values(["constraint_id", "status", "company"])
    table(wb, "Companies", co, [
        ("company", "Company", 24, "text"), ("constraint", "Builds on", 28, "text"), ("what_they_do", "What they do", 52, "text"),
        ("status", "Status", 11, "tag"), ("status_evidence_url", "Evidence", 13, "url"), ("consensus", "Well known?", 9, "text"),
        ("country", "Country", 13, "text"), ("founded", "Founded", 9, "text"), ("last_funding_usd", "Last round (USD)", 13, "usd"),
        ("last_funding_date", "Last round date", 11, "text"), ("note", "Notes", 60, "text")],
        intro="One row per company and constraint, so a company can appear more than once.")
    table(wb, "Funding rounds", fe.sort_values(["company", "date"]) if not fe.empty else fe, [
        ("company", "Company", 26, "text"), ("date", "Date", 11, "text"), ("amount_usd", "Amount (USD)", 14, "usd"),
        ("round", "Round", 18, "text"), ("source_url", "Link", 14, "url"), ("note", "Notes", 70, "text")],
        intro="Disclosed rounds in the activity window. Blank amount = not disclosed in USD.")
    if not act.empty:
        act = act.copy()
        act["constraint"] = act.constraint_id.map(lab)
    table(wb, "Activity", act, [
        ("constraint", "Constraint", 40, "text"), ("companies_count", "Companies", 10, "int"), ("funding_usd_5y", "Funding", 15, "usd"),
        ("launches_count", "Launches", 9, "int"), ("papers_last3y", "Papers, last 3 years", 10, "int"),
        ("papers_prev3y", "Papers, 3 years before", 10, "int"), ("notes", "Breakdown", 60, "text")],
        intro="Blank = nothing found, which means unknown, not zero.")
    table(wb, "Failures & exits", fail, [
        ("company", "Company", 26, "text"), ("year", "Year", 8, "text"), ("what_failed", "What happened", 60, "text"),
        ("reason", "Why", 60, "text"), ("source_url", "Link", 14, "url")])
    if not absn.empty:
        absn = absn.copy()
        absn["constraint"] = absn.constraint_id.map(lab)
    table(wb, "Searches, no result", absn, [
        ("constraint", "Constraint", 30, "text"), ("query", "Search", 50, "text"), ("date_searched", "Date", 11, "text"),
        ("result_summary", "What came back", 80, "text")],
        intro="Evidence of absence: searches that found nothing relevant.")
    intro = None
    if not ver.empty:
        ver = ver.copy()
        ver["verdict"] = ver.verdict.str.upper().str.strip()
        if "item" not in ver:
            ver["item"] = ""
        if "row_id" in ver:
            ver["item"] = ver.item.where(ver.item != "", ver.row_id)
        vc = ver.verdict.value_counts()
        n = len(ver)
        intro = (f"{vc.get('MATCH', 0)} of {n} checks matched ({vc.get('MATCH', 0) / n * 100:.1f}%), "
                 f"{vc.get('MISMATCH', 0)} mismatched (corrected before the report), {vc.get('UNVERIFIABLE', 0)} could not be opened.")
    table(wb, "Verification", ver, [
        ("item", "Row or item", 10, "text"), ("file", "Table", 14, "text"), ("field_checked", "What was checked", 20, "text"),
        ("verdict", "Result", 14, "tag"), ("claimed", "Claimed", 36, "text"), ("found", "Found on the page", 36, "text"),
        ("url_opened", "Link", 13, "url"), ("note", "Notes", 50, "text")], intro=intro)

    out = a.out or f"paradigma_{re.sub(r'[^a-z0-9]+', '-', (a.field or 'scan').lower()).strip('-')}.xlsx"
    wb.save(out)
    print(f"wrote {out}: {len(wb.worksheets)} sheets")


# -------------------------------------------------------------------- main
def main():
    p = argparse.ArgumentParser(description="PARADIGMA analysis engine")
    sp = p.add_subparsers(dest="cmd", required=True)

    c = sp.add_parser("cpi"); c.add_argument("--out", default="cpi.csv"); c.set_defaults(f=cmd_cpi)

    c = sp.add_parser("curves")
    c.add_argument("--points", default="points.csv")
    c.add_argument("--constraints", default="constraints.csv")
    c.add_argument("--cpi", default="cpi.csv")
    c.add_argument("--base-year", type=int)
    c.add_argument("--period", choices=["auto", "year", "quarter"], default="auto")
    c.add_argument("--min-n", type=int, default=3)
    c.add_argument("--start", type=int, help="first year of the analysis window (points before it are rejected)")
    c.add_argument("--end", type=int, help="last year of the analysis window (points after it are rejected)")
    c.add_argument("--outdir", default=".")
    c.add_argument("--plots", action="store_true")
    c.set_defaults(f=cmd_curves)

    c = sp.add_parser("gap")
    c.add_argument("--constraints", default="constraints.csv")
    c.add_argument("--curves", default="curves.csv")
    c.add_argument("--activity", default="activity.csv")
    c.add_argument("--cut", type=float, default=0.5)
    c.add_argument("--whynot", default="whynot.csv", help="colours the dots by why-not label when the file exists")
    c.add_argument("--title", help='chart title, e.g. "Gap matrix: dental aligners, 2016-2025"')
    c.add_argument("--outdir", default=".")
    c.add_argument("--plots", action="store_true")
    c.set_defaults(f=cmd_gap)

    c = sp.add_parser("papers")
    c.add_argument("--queries", default="paper_queries.csv")
    c.add_argument("--start", type=int, default=2005)
    c.add_argument("--end", type=int, default=2025)
    c.add_argument("--out", default="papers.csv")
    c.set_defaults(f=cmd_papers)

    c = sp.add_parser("workbook", help="readable Excel workbook from the run's CSVs (needs openpyxl)")
    c.add_argument("--field")
    c.add_argument("--region", default="global")
    c.add_argument("--start", type=int)
    c.add_argument("--end", type=int)
    c.add_argument("--date", help="run date, YYYY-MM-DD")
    c.add_argument("--constraints", default="constraints.csv")
    c.add_argument("--points", default="points.csv")
    c.add_argument("--out", help="output .xlsx (default paradigma_<field>.xlsx)")
    c.set_defaults(f=cmd_workbook)

    a = p.parse_args()
    if hasattr(a, "outdir"):
        os.makedirs(a.outdir, exist_ok=True)
    a.f(a)


if __name__ == "__main__":
    main()
